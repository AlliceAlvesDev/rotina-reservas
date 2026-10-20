import unittest
from io import BytesIO
import zipfile

import pandas as pd
from openpyxl import load_workbook
from reservas import comparar, excel, gerar_textos, normalizar_pedido, pacote, INTRODUCAO, MODELO_PEDIDO


class ReservasTest(unittest.TestCase):
    def test_mensagem_com_introducao_unica(self):
        dados = pd.DataFrame([["01/10/26 12:30", "12", "Cliente: Maria"], ["02/10/26", "13", "João"]])
        texto = gerar_textos(dados, MODELO_PEDIDO, INTRODUCAO)
        self.assertEqual(texto.count("Bom dia!"), 1)
        self.assertEqual(texto.count("Favor verificar"), 1)
        self.assertEqual(texto.count("Pedido: "), 2)
        self.assertIn("Pedido: 12\nData: 01/10/26 12:30\nCliente: Maria", texto)
        self.assertIn("Cliente: João", texto)
        self.assertNotIn("Cliente: Cliente:", texto)
        self.assertEqual(gerar_textos(dados.iloc[:0], MODELO_PEDIDO, INTRODUCAO), "")

    def test_normalizacao(self):
        for valor in [123, 123.0, " 001 23 ", "123.0"]:
            self.assertEqual(normalizar_pedido(valor), "123")
        self.assertEqual(normalizar_pedido(""), "")
        with self.assertRaises(ValueError):
            normalizar_pedido("PRODUTO")

    def test_comparacao_e_vazios(self):
        atual = pd.DataFrame([["primeira", "001"], ["repetida", "1.0"], ["outro", "2"], ["sem", ""], ["sem2", ""]])
        reserva = pd.DataFrame([["comum", "1"], ["exclusivo", "3"], ["repetida", "03"]])
        r, vazios = comparar(atual, reserva)
        self.assertEqual(vazios, 2)
        self.assertEqual(len(r["Atual sem duplicatas"]), 4)
        self.assertEqual(r["Somente na Atual"].iloc[0, 1], "2")
        self.assertEqual(r["Somente na Reserva"].iloc[0, 1], "3")
        self.assertEqual(r["Atual sem duplicatas"].iloc[0, 0], "primeira")
        self.assertEqual(gerar_textos(r["Somente na Reserva"], "Pedido {pedido}: {A}"), "Pedido 3: exclusivo")
        with self.assertRaises(ValueError):
            gerar_textos(reserva, "{campo_inexistente}")
        with zipfile.ZipFile(BytesIO(pacote(r, "texto"))) as z:
            self.assertEqual(len(z.namelist()), 4)

    def test_formula_exportada_como_texto(self):
        dados = pd.DataFrame({"a": ["=1+1"], "b": ["001"]})
        ws = load_workbook(BytesIO(excel({"dados": dados}))).active
        self.assertEqual(ws["A2"].data_type, "s")
        self.assertEqual(ws["B2"].value, "001")


if __name__ == "__main__":
    unittest.main()
